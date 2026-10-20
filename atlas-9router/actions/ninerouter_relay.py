"""
ask_9router — the relay between what Gemini hears/sees and what 9Router answers.

Gemini Live calls this on every user turn while the 9Router bridge is on (see
the [9ROUTER RELAY] block main.py adds to the system prompt). The handler
gathers the media the turn needs, sends it to 9Router through core/ninerouter.py
and returns the answer for Gemini to speak.

Everything 9Router receives is TEXT by default (ninerouter "text_only"): the
media gathered here is decoded by Gemini into explicit text and code before it
leaves — see core/ninerouter.decode_media.

Media, "refined raw" before it goes out:
    text    the request Gemini wrote + the raw transcript, de-stuttered
    screen  one screenshot (reuses screen_processor's capture + compression)
    camera  one webcam frame
    file    image  -> the image itself (downscaled)
            video  -> evenly spaced frames, because /chat/completions takes
                      images and not video
            pdf / docx / pptx / text / code -> the extracted text, capped
"""
from __future__ import annotations

import io
from pathlib import Path

from core import ninerouter

_IMG_EXT   = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
_VIDEO_EXT = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v", ".wmv"}
_AUDIO_EXT = {".mp3", ".wav", ".ogg", ".m4a", ".flac", ".aac", ".opus"}
_AUDIO_MIME = {".mp3": "audio/mp3", ".wav": "audio/wav", ".ogg": "audio/ogg",
               ".m4a": "audio/mp4", ".flac": "audio/flac", ".aac": "audio/aac",
               ".opus": "audio/ogg"}
_TEXT_CAP  = 20000
_VIDEO_FRAMES = 6
_MAX_SIDE  = 1280


def _shrink(img_bytes: bytes, mime: str) -> tuple[bytes, str]:
    """Downscale big images; 9Router providers bill and time out by size."""
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(img_bytes))
        if max(im.size) <= _MAX_SIDE and mime in ("image/jpeg", "image/png"):
            return img_bytes, mime
        im = im.convert("RGB")
        im.thumbnail((_MAX_SIDE, _MAX_SIDE))
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=82)
        return buf.getvalue(), "image/jpeg"
    except Exception:
        return img_bytes, mime


def _video_frames(path: Path) -> list:
    import cv2
    cap = cv2.VideoCapture(str(path))
    total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    fps = cap.get(cv2.CAP_PROP_FPS) or 0
    out = []
    if total <= 0:
        cap.release()
        return out
    for i in range(_VIDEO_FRAMES):
        idx = int(total * (i + 0.5) / _VIDEO_FRAMES)
        cap.set(cv2.CAP_PROP_POS_FRAMES, idx)
        ok, frame = cap.read()
        if not ok or frame is None:
            continue
        ok, jpg = cv2.imencode(".jpg", frame, [cv2.IMWRITE_JPEG_QUALITY, 80])
        if ok:
            b, m = _shrink(jpg.tobytes(), "image/jpeg")
            t = f"{idx / fps:.1f}s" if fps else f"frame {idx}"
            out.append((b, m, f"video frame {i + 1}/{_VIDEO_FRAMES} at {t}"))
    cap.release()
    return out


def _doc_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            try:
                import pdfplumber
                with pdfplumber.open(str(path)) as pdf:
                    return "\n".join((p.extract_text() or "") for p in pdf.pages[:40])
            except ImportError:
                from PyPDF2 import PdfReader
                return "\n".join((p.extract_text() or "")
                                 for p in PdfReader(str(path)).pages[:40])
        if ext == ".docx":
            import docx
            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        if ext == ".pptx":
            from pptx import Presentation
            parts = []
            for n, slide in enumerate(Presentation(str(path)).slides, 1):
                parts.append(f"--- slide {n} ---")
                for sh in slide.shapes:
                    if getattr(sh, "has_text_frame", False):
                        parts.append(sh.text_frame.text)
            return "\n".join(parts)
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"(could not read {path.name}: {e})"


def _gather(attach: str, file_path: str, player) -> tuple[list, str, str]:
    """-> (images, extra_text, note for the log)"""
    images, extra, notes = [], "", []
    attach = (attach or "none").lower().strip()

    if attach in ("screen", "both"):
        from actions.screen_processor import _capture_screen
        b, m = _capture_screen()
        images.append((*_shrink(b, m), "screen capture — the user's COMPUTER"))
        notes.append("screen")
    if attach in ("camera", "both"):
        from actions.screen_processor import _capture_camera
        b, m = _capture_camera()
        images.append((*_shrink(b, m), "webcam frame — the USER"))
        notes.append("camera")

    if attach == "file" or file_path:
        fp = file_path or (getattr(player, "current_file", None) or "")
        p = Path(str(fp)).expanduser() if fp else None
        if not p or not p.exists():
            extra = f"(the user referred to a file, but none was found at '{fp}')"
        else:
            ext = p.suffix.lower()
            if ext in _IMG_EXT:
                images.append((*_shrink(p.read_bytes(),
                                        "image/png" if ext == ".png" else "image/jpeg"),
                               f"image file {p.name}"))
            elif ext in _AUDIO_EXT:
                # Never sent as audio: core.ninerouter decodes it to a transcript.
                images.append((p.read_bytes(), _AUDIO_MIME.get(ext, "audio/mpeg"),
                               f"audio file {p.name}"))
            elif ext in _VIDEO_EXT:
                frames = _video_frames(p)
                images.extend(frames)
                extra = f"Video file {p.name}: {len(frames)} frames attached in order."
            else:
                txt = _doc_text(p)
                if len(txt) > _TEXT_CAP:
                    txt = txt[:_TEXT_CAP] + f"\n…(truncated, {len(txt)} chars total)"
                extra = f"File {p.name}:\n{txt}"
            notes.append(f"file:{p.name}")
    return images, extra, ",".join(notes) or "text"


def ask_9router(parameters: dict, player=None) -> str:
    if not ninerouter.enabled():
        return ("The 9Router bridge is switched off. Answer the user yourself "
                "and do not call ask_9router again this session.")

    request = ninerouter.refine(str(parameters.get("request") or ""))
    raw = ninerouter.last_heard(consume=True) or ""
    if not request:
        request = ninerouter.refine(raw)
    if not request:
        return "Nothing to send: the request was empty. Ask the user to repeat."

    try:
        images, extra, _note = _gather(parameters.get("attach", "none"),
                                       str(parameters.get("file_path") or ""),
                                       player)
    except Exception as e:
        images, extra = [], f"(attachment failed: {e})"

    try:
        answer = ninerouter.ask(request, raw_heard=raw, images=images,
                                extra_text=extra)
    except Exception as e:
        return (f"[9ROUTER FAILED] {e}. Tell the user in one short sentence that "
                f"the 9Router brain did not answer and why, then answer the "
                f"question yourself as best you can, saying it is a fallback.")

    try:
        if player is not None and len(answer) > 280:
            model = ninerouter.stats.get("last_model") or "9router"
            player.show_content(f"9ROUTER · {model}", answer)
    except Exception:
        pass

    return ("[9ROUTER ANSWER] Speak the text below to the user faithfully, in "
            "their language. Do not add facts, do not shorten meaning, do not "
            "mention 9Router unless asked. Skip code blocks aloud (say it is on "
            "screen).\n\n" + answer)


TOOL = {
    "name": "ask_9router",
    "description": (
        "Sends the user's request to the 9Router reasoning model and returns "
        "the answer you must speak. While the 9Router bridge is active, call "
        "this for EVERY question, conversation turn or request for information, "
        "explanation, opinion, analysis, writing or code — never answer those "
        "from your own knowledge. Put a cleaned-up but faithful version of what "
        "the user said in 'request' (keep their language, fix only obvious "
        "transcription errors, add context from earlier turns if the request "
        "depends on it). Set 'attach' to 'screen' when the request is about "
        "what is on the computer, 'camera' when it is about the user or their "
        "room, 'both' when both matter, 'file' when it is about a dropped or "
        "named file (image, video, audio, PDF, document, code). Computer actions "
        "(open apps, volume, reminders, etc.) still use their own tools."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "request": {
                "type": "STRING",
                "description": "The user's request, refined but faithful, in their language.",
            },
            "attach": {
                "type": "STRING",
                "description": "none | screen | camera | both | file",
            },
            "file_path": {
                "type": "STRING",
                "description": "Path of the file when attach='file'. Empty = the file dropped in the UI.",
            },
        },
        "required": ["request"],
    },
    "handler": ask_9router,
    "tier": "core",
    "brief": "Relay the user's request (and optional screen/camera/file) to 9Router and speak its answer.",
}
